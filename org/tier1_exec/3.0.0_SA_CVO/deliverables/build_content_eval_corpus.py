#!/usr/bin/env python3
"""Build the ~100-doc labeled eval corpus from the LOCAL SharePoint sync (Andrew-direct via CTO~Z).

Reads real docs off the local OneDrive sync (no gateway/base64), extracts text (pdf/docx/pptx/xlsx),
writes one .txt per doc + a manifest to C:/Users/AndrewBerman/content_eval/ (OUTSIDE the repo -- raw
document text NEVER enters git). The answer_key is built SEPARATELY by CVO judgment (not by this script,
not by any candidate axis).

Storage-care: output dir is outside the repo. This SCRIPT (code, no client text) lives in the repo.
"""
from __future__ import annotations

import datetime as _dt
import json
import os
import re
import sys

SYNC = "C:/Users/AndrewBerman/OneDrive - Opal Capital/Shared - Documents"
OUT = "C:/Users/AndrewBerman/content_eval"
DOCS = os.path.join(OUT, "docs")
MAXCHARS = 12000

# Family-rich + topically-diverse folders (cap per folder for variety); "" = loose root files.
PLAN = [
    ("", 13),                    # loose root: factsheets/decks/portfolios/strategy -- cross-topic
    ("49F Communication", 14),   # project set (+ likely version pairs)
    ("Market Update", 16),       # DATED SERIES -> timing relationships
    ("Div Research", 12),        # research series
    ("Presentations", 10),       # decks
    ("Investment Committee", 8),
    ("Fund Comps", 6),
    ("Macro Research", 4),
    ("Compliance Manual", 5),
    ("Marketing", 8),
    ("Viewpoints", 6),
]
EXTS = (".pdf", ".docx", ".pptx", ".xlsx")
SKIP_BIGGER_THAN = 25 * 1024 * 1024  # skip >25MB (huge decks) to keep extraction fast


def _text_pdf(p):
    import fitz
    doc = fitz.open(p)
    out = []
    for i, page in enumerate(doc):
        if i >= 40:
            break
        out.append(page.get_text())
    doc.close()
    return "\n".join(out)


def _text_docx(p):
    import docx
    d = docx.Document(p)
    parts = [para.text for para in d.paragraphs]
    for tbl in d.tables:
        for row in tbl.rows:
            parts.append(" ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _text_pptx(p):
    from pptx import Presentation
    prs = Presentation(p)
    parts = []
    for i, slide in enumerate(prs.slides):
        if i >= 80:
            break
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def _text_xlsx(p):
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append("# sheet: %s" % ws.title)
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            if r >= 200:
                break
            parts.append(" ".join("" if c is None else str(c) for c in row))
    wb.close()
    return "\n".join(parts)


EXTRACT = {".pdf": _text_pdf, ".docx": _text_docx, ".pptx": _text_pptx, ".xlsx": _text_xlsx}


def _slug(name):
    return re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_")[:60]


def _list_folder(folder):
    base = os.path.join(SYNC, folder) if folder else SYNC
    out = []
    try:
        for nm in sorted(os.listdir(base)):
            fp = os.path.join(base, nm)
            if not os.path.isfile(fp):
                continue
            ext = os.path.splitext(nm)[1].lower()
            if ext not in EXTS:
                continue
            try:
                sz = os.path.getsize(fp)
            except OSError:
                continue
            if sz == 0 or sz > SKIP_BIGGER_THAN:
                continue
            out.append((nm, fp, ext, sz))
    except OSError:
        pass
    return out


def main():
    os.makedirs(DOCS, exist_ok=True)
    manifest, did = [], 0
    ftmix = {}
    for folder, cap in PLAN:
        files = _list_folder(folder)
        # format-balance within a folder: interleave by ext, prefer smaller within each ext
        by_ext = {}
        for f in files:
            by_ext.setdefault(f[2], []).append(f)
        for e in by_ext:
            by_ext[e].sort(key=lambda x: x[3])
        picked, exts_cycle = [], sorted(by_ext)
        idx = {e: 0 for e in exts_cycle}
        while len(picked) < cap and any(idx[e] < len(by_ext[e]) for e in exts_cycle):
            for e in exts_cycle:
                if idx[e] < len(by_ext[e]) and len(picked) < cap:
                    picked.append(by_ext[e][idx[e]])
                    idx[e] += 1
        for nm, fp, ext, sz in picked:
            try:
                text = EXTRACT[ext](fp)[:MAXCHARS]
            except Exception as ex:  # noqa: BLE001
                print("SKIP (extract fail): %s (%s)" % (nm, ex))
                continue
            if len(text.strip()) < 40:
                print("SKIP (near-empty text): %s" % nm)
                continue
            did += 1
            doc_id = "d%03d" % did
            with open(os.path.join(DOCS, doc_id + ".txt"), "w", encoding="utf-8") as fh:
                fh.write(text)
            try:
                mtime = _dt.datetime.utcfromtimestamp(os.path.getmtime(fp)).isoformat() + "Z"
            except OSError:
                mtime = None
            manifest.append({"doc_id": doc_id, "name": nm, "folder": folder or "(root)",
                             "source_relpath": os.path.relpath(fp, SYNC).replace("\\", "/"),
                             "filetype": ext.lstrip("."), "size": sz, "mtime_utc": mtime,
                             "chars": len(text), "author": None})
            ftmix[ext] = ftmix.get(ext, 0) + 1
    with open(os.path.join(OUT, "manifest.json"), "w", encoding="utf-8") as fh:
        json.dump({"generated_note": "local SharePoint sync extract; author=None (needs Graph); mtime=disk",
                   "n_docs": len(manifest), "filetype_mix": {k.lstrip('.'): v for k, v in ftmix.items()},
                   "docs": manifest}, fh, indent=1)
    print("DONE: %d docs -> %s" % (len(manifest), DOCS))
    print("filetype mix:", {k.lstrip('.'): v for k, v in ftmix.items()})
    print("by folder:", {})
    fold = {}
    for m in manifest:
        fold[m["folder"]] = fold.get(m["folder"], 0) + 1
    for k, v in fold.items():
        print("  %2d  %s" % (v, k))
    return 0


if __name__ == "__main__":
    sys.exit(main())
